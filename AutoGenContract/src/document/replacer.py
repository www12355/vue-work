"""python-docx 替换引擎。

直接操作 OOXML，不依赖 win32com / Microsoft Word。
在 XML 文本节点层面做正则替换，无 \\r \\n 控制字符问题。

用法：
    replacer = DocxReplacer(template_path, output_path)
    replacer.replace_all(values_dict)
    replacer.save()
"""

import re
import os
import logging
from copy import deepcopy
from typing import Dict, List, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)

# 匹配 {{ name }} 或 {{name}}
PLACEHOLDER_RE = re.compile(r"\{\{\s*(\w+)\s*\}\}")

LLM_FIELDS = {
    "aicc_service_content",
    "aicc_searcher_protected_content",
    "aicc_provider_protected_content",
    "aicc_service_fee",
    "aicc_delete_provide_condition",
}


class DocxReplacer:
    """python-docx 替换引擎。只读模板，写入新文件。"""

    def __init__(self, template_path: str, output_path: str):
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"模板文件不存在: {template_path}")
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        self.doc = Document(template_path)
        self.output_path = output_path

    # ========== 公共接口 ==========

    def replace_all(self, values: Dict[str, str]):
        """替换所有占位符。简单字段用正则替换，LLM 字段用多段落替换。"""
        simple = {}
        llm = {}
        for k, v in values.items():
            (llm if k in LLM_FIELDS else simple)[k] = str(v) if v else ""

        logger.info(f"替换: {len(simple)} 简单 + {len(llm)} LLM")
        self._replace_in_paragraphs(simple)
        self._replace_in_tables(simple)
        logger.info(f"简单字段完成 ({len(simple)} 个)")
        for name, text in llm.items():
            self._replace_llm_section(name, text)
        logger.info("全部替换完成")

    def save(self):
        self.doc.save(self.output_path)
        logger.info(f"已保存: {self.output_path}")

    # ========== 简单字段 ==========

    def _replace_in_paragraphs(self, simple: Dict[str, str]):
        if not simple:
            return
        count = 0
        for para in self.doc.paragraphs:
            count += self._replace_in_para(para, simple)
        logger.info(f"段落: {count} 次替换")

    def _replace_in_tables(self, simple: Dict[str, str]):
        if not simple:
            return
        count = 0
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        count += self._replace_in_para(para, simple)
        logger.info(f"表格: {count} 次替换")

    def _replace_in_para(self, para, simple: Dict[str, str]) -> int:
        """在一个段落的所有 run 中做正则替换。"""
        count = 0
        for run in para.runs:
            text = run.text
            new_text = text
            for name, value in simple.items():
                ptn = re.compile(r"\{\{\s*" + re.escape(name) + r"\s*\}\}")
                new_text, n = ptn.subn(value, new_text)
                count += n
            if new_text != text:
                run.text = new_text
        return count

    # ========== LLM 多段落 ==========

    def _replace_llm_section(self, name: str, text: str):
        """找到含占位符的段落，替换为 LLM 生成的多段落文本。"""
        if not text or not text.strip():
            logger.warning(f"LLM 字段 {name} 为空，跳过")
            return

        lines = [l for l in text.strip().split("\n") if l.strip()]
        if not lines:
            return

        # 在段落中查找
        for para in self.doc.paragraphs:
            if not self._para_contains(para, name):
                continue

            # 清空当前段落内容，填入第一行
            self._clear_para(para)
            para.add_run(lines[0])

            # 在当前段落后插入新段落
            last_elem = para._element
            for line in lines[1:]:
                last_elem = self._insert_para_after(last_elem, line)

            logger.info(f"LLM: {name} → {len(lines)} 段")
            return

        # 在表格中查找
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if self._para_contains(para, name):
                            self._clear_para(para)
                            para.add_run(lines[0])
                            logger.info(f"LLM(表格): {name}")
                            return

        logger.warning(f"LLM 占位符未找到: {name}")

    # ========== 辅助 ==========

    def _para_contains(self, para, placeholder_name: str) -> bool:
        """检查段落是否包含指定占位符。"""
        return bool(PLACEHOLDER_RE.search(para.text)) and \
               any(m.group(1) == placeholder_name
                   for m in PLACEHOLDER_RE.finditer(para.text))

    @staticmethod
    def _clear_para(para):
        """清空段落的所有内容（run、图片等）。"""
        runs = para.runs
        for r in runs:
            r._element.getparent().remove(r._element)

    @staticmethod
    def _insert_para_after(elem, text: str):
        """在给定 XML 元素后插入一个新段落，返回新段落的元素。"""
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        p.append(r)
        elem.addnext(p)
        return p
